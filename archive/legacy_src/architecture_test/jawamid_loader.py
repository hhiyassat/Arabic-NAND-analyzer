"""Load الجوامد classification table for architecture tests.

Computational-level data loader. Bridges data/الجوامد.xlsx with the
linguistic level (per 11_Abstraction_Levels.md). Provides per-token jamid
category lookup using the seed examples in the classification table.

NOT a closed lexicon — the table is a *seed*, not an exhaustive list.
Negative result == "not in seed", NOT "is not jamid".
"""

from __future__ import annotations

import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

_REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_JAWAMID_XLSX = _REPO_ROOT / "data" / "الجوامد.xlsx"

_DIACRITICS = set("ًٌٍَُِّْٰٓٔ")


def _nfc(text: str) -> str:
    return unicodedata.normalize("NFC", str(text or "").strip())


def _strip_diacritics(text: str) -> str:
    d = unicodedata.normalize("NFD", _nfc(text))
    return unicodedata.normalize(
        "NFC",
        "".join(c for c in d if c not in _DIACRITICS and unicodedata.category(c) != "Mn"),
    )


def _normalize_lookup_key(text: str) -> str:
    """Aggressive normalization for jamid lookup: strip diacritics + hamza variants.

    Also strips a trailing alif that came from tanwin-fath (يَوْمًا → يوم).
    """
    raw = str(text or "")
    has_tanwin = any(c in raw for c in "ًٌٍ")
    s = _strip_diacritics(raw)
    s = s.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
    s = s.replace("ى", "ي").replace("ة", "ه")
    s = s.strip()
    # Remove trailing alif left over from tanwin-fath.
    if has_tanwin and s.endswith("ا") and len(s) > 3:
        s = s[:-1]
    return s


@dataclass(frozen=True)
class JawamidCategory:
    serial: int
    category_name: str
    definition: str
    examples: tuple[str, ...]
    notes: str
    requires_root_pattern: str  # هل يُطلب له جذر/وزن؟


@dataclass
class JawamidIndex:
    """In-memory index for jamid noun classification."""

    categories: list[JawamidCategory] = field(default_factory=list)
    _example_to_categories: dict[str, list[str]] = field(default_factory=dict)

    def __post_init__(self) -> None:
        for cat in self.categories:
            for ex in cat.examples:
                key = _normalize_lookup_key(ex)
                if not key:
                    continue
                self._example_to_categories.setdefault(key, []).append(cat.category_name)

    def lookup(self, surface: str) -> list[str]:
        """Return jamid categories matching the surface form (could be multiple).

        Empty list means "not in seed examples" — NOT a negative classification.
        Some words intentionally appear in multiple categories (e.g., "من" is both
        اسم موصول and اسم شرط and اسم استفهام).
        """
        key = _normalize_lookup_key(surface)
        if not key:
            return []
        return list(self._example_to_categories.get(key, []))

    def category_count(self) -> int:
        return len(self.categories)

    def example_count(self) -> int:
        return sum(len(c.examples) for c in self.categories)


def load_jawamid_index(path: Path | None = None) -> JawamidIndex:
    xlsx = path or DEFAULT_JAWAMID_XLSX
    if not xlsx.is_file():
        raise FileNotFoundError(f"Jawamid dataset not found: {xlsx}")

    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["Jawamid_Classification"]

    categories: list[JawamidCategory] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        serial_raw = row[0]
        try:
            serial = int(serial_raw)
        except (TypeError, ValueError):
            continue
        examples_str = str(row[3] or "")
        # Examples are comma-separated (Arabic and ASCII commas).
        raw_examples = [
            e.strip()
            for e in examples_str.replace("،", ",").replace("؛", ",").split(",")
        ]
        examples = tuple(e for e in raw_examples if e)
        categories.append(
            JawamidCategory(
                serial=serial,
                category_name=str(row[1] or "").strip(),
                definition=str(row[2] or "").strip(),
                examples=examples,
                notes=str(row[4] or "").strip(),
                requires_root_pattern=str(row[5] or "").strip(),
            )
        )
    return JawamidIndex(categories=categories)
