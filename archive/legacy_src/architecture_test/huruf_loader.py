"""Load Huruf unified master dataset for architecture tests."""

from __future__ import annotations

import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Optional

_REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_HURUF_XLSX = (
    _REPO_ROOT
    / "huruf_maani_normalization"
    / "huruf_maani_unified_master.xlsx"
)

_DIACRITICS = set("ًٌٍَُِّْٰٓٔ")


def strip_diacritics(text: str) -> str:
    d = unicodedata.normalize("NFD", str(text or ""))
    return unicodedata.normalize(
        "NFC",
        "".join(c for c in d if c not in _DIACRITICS and unicodedata.category(c) != "Mn"),
    )


def nfc(text: str) -> str:
    return unicodedata.normalize("NFC", str(text or "").strip())


@dataclass(frozen=True)
class HurufEntry:
    canonical_id: str
    canonical_form: str
    fully_vocalized_form: str
    normalized_form: str
    category: str
    functional_roles: str
    linguistic_classification: str
    notes: str


class HurufIndex:
    def __init__(self, entries: list[HurufEntry]) -> None:
        self.entries = entries
        self._by_id = {e.canonical_id: e for e in entries}
        self._by_vocalized: dict[str, HurufEntry] = {}
        self._by_plain: dict[str, list[HurufEntry]] = {}
        for e in entries:
            for key in (e.fully_vocalized_form, e.canonical_form, e.normalized_form):
                plain = strip_diacritics(nfc(key))
                if plain:
                    self._by_plain.setdefault(plain, []).append(e)
            voc = nfc(e.fully_vocalized_form)
            if voc:
                self._by_vocalized[voc] = e

    def lookup(self, surface: str, *, expected_id: str | None = None) -> Optional[HurufEntry]:
        if expected_id and expected_id in self._by_id:
            return self._by_id[expected_id]
        voc = nfc(surface)
        if voc in self._by_vocalized:
            return self._by_vocalized[voc]
        plain = strip_diacritics(voc)
        hits = self._by_plain.get(plain, [])
        if not hits:
            return None
        if len(hits) == 1:
            return hits[0]
        if expected_id:
            for h in hits:
                if h.canonical_id == expected_id:
                    return h
        return hits[0]

    def match_report(self, surface: str, *, expected_id: str | None = None) -> dict[str, Any]:
        entry = self.lookup(surface, expected_id=expected_id)
        return {
            "surface": surface,
            "in_dataset": entry is not None,
            "canonical_id": entry.canonical_id if entry else None,
            "canonical_form": entry.canonical_form if entry else None,
            "fully_vocalized_form": entry.fully_vocalized_form if entry else None,
            "category": entry.category if entry else None,
            "functional_roles": entry.functional_roles if entry else None,
        }


def load_huruf_index(path: Path | None = None) -> HurufIndex:
    xlsx = path or DEFAULT_HURUF_XLSX
    if not xlsx.is_file():
        raise FileNotFoundError(f"Huruf dataset not found: {xlsx}")

    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["Unified_Master"]
    entries: list[HurufEntry] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        entries.append(
            HurufEntry(
                canonical_id=str(row[0]).strip(),
                canonical_form=str(row[1] or "").strip(),
                fully_vocalized_form=str(row[2] or "").strip(),
                normalized_form=str(row[5] or "").strip(),
                category=str(row[6] or "").strip(),
                functional_roles=str(row[8] or "").strip(),
                linguistic_classification=str(row[9] or "").strip(),
                notes=str(row[13] or "").strip(),
            )
        )
    return HurufIndex(entries)
