"""Load Mushtaqat (derived nouns) weights table for architecture tests.

Computational-level data loader. Bridges
data/Mushtaqat_Weights_Final_Corrected_With_Fa3l-minimall.xlsx with the
linguistic level (per 11_Abstraction_Levels.md). Each row is a (وزن، باب)
pair with example forms; we expose pattern-based matching from token surface
to derivational category.

Pattern matching is a HEURISTIC. The table is reference, not a parser.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

_REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_MUSHTAQAT_XLSX = (
    _REPO_ROOT / "data" / "Mushtaqat_Weights_Final_Corrected_With_Fa3l-minimall.xlsx"
)

_DIACRITICS = set("ًٌٍَُِّْٰٓٔ")


def _nfc(text: str) -> str:
    return unicodedata.normalize("NFC", str(text or "").strip())


def _strip_diacritics(text: str) -> str:
    d = unicodedata.normalize("NFD", _nfc(text))
    return unicodedata.normalize(
        "NFC",
        "".join(c for c in d if c not in _DIACRITICS and unicodedata.category(c) != "Mn"),
    )


def _consonant_skeleton(text: str) -> str:
    """Extract the consonant skeleton (no diacritics, no shadda)."""
    s = _strip_diacritics(text)
    # Keep only Arabic consonants and basic letters
    return "".join(c for c in s if "؀" <= c <= "ۿ" and c not in "ـ")


@dataclass(frozen=True)
class MushtaqWeight:
    serial: int
    bab: str  # باب المشتق (اسم الفاعل، اسم المفعول، الصفة المشبهة...)
    wazn_voweled: str  # الوزن مضبوطًا (e.g., فَاعِل، مُفْعِل)
    feminine_form: str  # الصورة بالتاء / المؤنث
    example_voweled: str  # مثال مضبوط
    example_feminine: str  # مثال مؤنث مضبوط
    classification: str  # قياسي/سماعي/محدث
    is_productive: str  # منتج؟
    accepts_ta: str  # يقبل التاء؟
    note: str


@dataclass
class MushtaqatIndex:
    """In-memory index for derived noun pattern matching."""

    weights: list[MushtaqWeight] = field(default_factory=list)
    _by_skeleton: dict[str, list[MushtaqWeight]] = field(default_factory=dict)
    _by_bab: dict[str, list[MushtaqWeight]] = field(default_factory=dict)

    def __post_init__(self) -> None:
        for w in self.weights:
            # Index by skeleton of example
            ex_skel = _consonant_skeleton(w.example_voweled)
            if ex_skel:
                self._by_skeleton.setdefault(ex_skel, []).append(w)
            # Index by بَاب
            self._by_bab.setdefault(w.bab, []).append(w)

    def lookup_by_example(self, surface: str) -> list[MushtaqWeight]:
        """Return all weights whose example matches surface (skeleton match)."""
        skel = _consonant_skeleton(surface)
        return list(self._by_skeleton.get(skel, []))

    def lookup_by_pattern_match(self, surface: str) -> list[MushtaqWeight]:
        """Match surface against weight pattern templates.

        Heuristic: builds regex from wazn's consonant pattern (ف ع ل),
        and checks if surface fits. Limited to patterns where the wazn
        is explicit (most rows). False positives expected — this is a
        reference signal, not a parser.
        """
        plain = _strip_diacritics(surface)
        if len(plain) < 3:
            return []
        matches = []
        for w in self.weights:
            if self._pattern_matches(w.wazn_voweled, plain):
                matches.append(w)
        return matches

    @staticmethod
    def _pattern_matches(wazn: str, surface_plain: str) -> bool:
        """Crude wazn → regex match. Treats ف/ع/ل as wildcard root slots.

        Examples:
          فَاعِل + كَاتِب → "فاعل" + "كاتب" → match (3 root letters)
          مُفْعِل + مُكْرِم → "مفعل" + "مكرم" → match (م + 3 root letters)
        """
        w_plain = _strip_diacritics(wazn)
        if not w_plain:
            return False
        # Build regex: ف/ع/ل become wildcards (Arabic letter class)
        # Other letters are literal.
        slots = []
        for ch in w_plain:
            if ch in "فعل":
                slots.append("[ء-ي]")
            else:
                slots.append(re.escape(ch))
        pattern = "^" + "".join(slots) + "$"
        try:
            return bool(re.match(pattern, surface_plain))
        except re.error:
            return False

    def weight_count(self) -> int:
        return len(self.weights)

    def bab_count(self) -> int:
        return len(self._by_bab)

    def babs(self) -> list[str]:
        return sorted(self._by_bab.keys())


def load_mushtaqat_index(path: Path | None = None) -> MushtaqatIndex:
    xlsx = path or DEFAULT_MUSHTAQAT_XLSX
    if not xlsx.is_file():
        raise FileNotFoundError(f"Mushtaqat dataset not found: {xlsx}")

    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["الأوزان_المصححة"]

    weights: list[MushtaqWeight] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        try:
            serial = int(row[0])
        except (TypeError, ValueError):
            continue
        weights.append(
            MushtaqWeight(
                serial=serial,
                bab=str(row[1] or "").strip(),
                wazn_voweled=str(row[2] or "").strip(),
                feminine_form=str(row[3] or "").strip(),
                example_voweled=str(row[4] or "").strip(),
                example_feminine=str(row[5] or "").strip(),
                classification=str(row[6] or "").strip(),
                is_productive=str(row[7] or "").strip(),
                accepts_ta=str(row[8] or "").strip(),
                note=str(row[9] or "").strip(),
            )
        )
    return MushtaqatIndex(weights=weights)
